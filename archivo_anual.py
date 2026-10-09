# archivo_anual.py
# ### CAMBIO AQUÍ ### — Archivo nuevo: ARCHIVO ANUAL descargable.
#
# Desde el 1° de enero, el Dashboard muestra un aviso fijo hasta que se descarga el archivo del año
# que terminó. El .zip trae:
#   1. datos_AAAA/     → reservas, reseñas, consultas e informes de ese año (CSV, se abren con Excel)
#   2. Resumen_AAAA.xlsx → métricas del año por mes, por propiedad y por moneda
#   3. base_completa/  → TODAS las tablas del sistema (todos los años) en CSV
#   4. codigo/         → el código del sistema (sin claves: las claves están solo en Render)
#   5. LEEME.txt       → qué hay en el archivo y cómo restaurar o arreglar el sistema
# Los datos NO se borran del sistema: se siguen usando para comparar años, precios e informes.

import csv
import io
import os
import zipfile
from datetime import date, datetime, timedelta, timezone

import database as db

AR = timezone(timedelta(hours=-3))
TABLAS_EXCLUIDAS = {"login_intentos"}
RAIZ = os.path.dirname(os.path.abspath(__file__))

# Código que se incluye (lista blanca: nunca se cuela un archivo de claves)
EXTENSIONES_CODIGO = {".py", ".html", ".css", ".js", ".txt", ".yml", ".yaml", ".md", ".png", ".svg", ".ico"}
NOMBRES_CODIGO = {"Procfile", ".python-version", ".env.example", "requirements.txt"}
CARPETAS_EXCLUIDAS = {"__pycache__", ".git", "venv", ".venv", "node_modules", "env"}


def hoy():
    return datetime.now(AR).date()


def _clave(anio):
    return f"archivo_anual_descargado_{anio}"


def anios_con_datos():
    """Años que tienen reservas cargadas (más nuevo primero)."""
    filas = db.query("""
        SELECT DISTINCT SUBSTRING(fecha_inicio FROM 1 FOR 4) AS a FROM reservas
        WHERE fecha_inicio ~ '^[0-9]{4}-' ORDER BY a DESC
    """)
    return [int(f["a"]) for f in filas if f["a"] and f["a"].isdigit()]


def descargado(anio):
    row = db.query_one("SELECT valor FROM config WHERE clave=%s", (_clave(anio),))
    return row["valor"] if row else None


def marcar_descargado(anio):
    db.execute("""
        INSERT INTO config (clave, valor) VALUES (%s, %s)
        ON CONFLICT (clave) DO UPDATE SET valor = EXCLUDED.valor
    """, (_clave(anio), datetime.now(AR).strftime("%Y-%m-%d %H:%M")))


def anio_pendiente():
    """Año terminado cuyo archivo todavía no se descargó (o None). Mira hasta 3 años atrás."""
    actual = hoy().year
    con_datos = set(anios_con_datos())
    for anio in range(actual - 1, actual - 4, -1):
        if anio in con_datos and not descargado(anio):
            return anio
    return None


# ── Contenido del zip ─────────────────────────────────────────────────────────
def _csv(filas):
    txt = io.StringIO()
    if filas:
        w = csv.DictWriter(txt, fieldnames=list(filas[0].keys()))
        w.writeheader()
        w.writerows(filas)
    return "﻿" + txt.getvalue()   # ﻿: Excel lee bien los acentos


def tablas():
    return [r["table_name"] for r in db.query("""
        SELECT table_name FROM information_schema.tables
        WHERE table_schema = 'public' AND table_type = 'BASE TABLE' ORDER BY table_name
    """) if r["table_name"] not in TABLAS_EXCLUIDAS]


def escribir_base_completa(z, carpeta=""):
    """Todas las tablas en CSV dentro del zip (la usa también 'Descargar copia')."""
    for t in tablas():
        filas = db.query(f'SELECT * FROM "{t}" ORDER BY 1')   # nombre sacado de la propia base
        z.writestr(f"{carpeta}{t}.csv", _csv(filas))


def _datos_del_anio(anio):
    desde, hasta = f"{anio}-01-01", f"{anio + 1}-01-01"
    return {
        "reservas": db.query("""
            SELECT r.*, p.nombre AS propiedad FROM reservas r LEFT JOIN propiedades p ON p.id = r.propiedad_id
            WHERE r.fecha_inicio >= %s AND r.fecha_inicio < %s ORDER BY r.fecha_inicio""", (desde, hasta)),
        "resenas": db.query("""
            SELECT r.*, p.nombre AS propiedad FROM resenas r LEFT JOIN propiedades p ON p.id = r.propiedad_id
            WHERE COALESCE(r.fecha,'') >= %s AND COALESCE(r.fecha,'') < %s ORDER BY r.fecha""", (desde, hasta)),
        "consultas_huespedes": db.query("""
            SELECT c.*, p.nombre AS propiedad FROM consultas_huespedes c LEFT JOIN propiedades p ON p.id = c.propiedad_id
            WHERE EXTRACT(YEAR FROM c.creada_en) = %s ORDER BY c.creada_en""", (anio,)),
        "informes": db.query("""
            SELECT * FROM informes WHERE EXTRACT(YEAR FROM creado_en) = %s ORDER BY creado_en""", (anio,)),
        "propiedades": db.query("SELECT * FROM propiedades ORDER BY id"),
    }


def _resumen_excel(anio, reservas):
    """Planilla con el resumen del año: por mes y por propiedad, separado por moneda."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    validas = [r for r in reservas if (r.get("estado") or "") != "cancelada"]
    canceladas = len(reservas) - len(validas)

    def agrupar(clave):
        g = {}
        for r in validas:
            k = (clave(r), r.get("moneda") or "ARS")
            d = g.setdefault(k, {"reservas": 0, "noches": 0, "bruto": 0.0, "comision": 0.0, "neto": 0.0})
            d["reservas"] += 1
            d["noches"] += int(r.get("noches") or 0)
            d["bruto"] += float(r.get("ingreso_bruto") or 0)
            d["comision"] += float(r.get("comision_airbnb") or 0)
            d["neto"] += float(r.get("ingreso_neto") or 0)
        return g

    wb = Workbook()
    negrita, fondo = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2A2A1E")
    cab = ["", "Moneda", "Reservas", "Noches", "Ingreso bruto", "Comisiones", "Ingreso neto", "Precio prom./noche"]

    def hoja(ws, titulo, filas_dict, etiqueta):
        ws.append([titulo]); ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Generado el {hoy().isoformat()} · reservas canceladas excluidas ({canceladas})"])
        ws.append([])
        ws.append([etiqueta] + cab[1:])
        for c in ws[4]:
            c.font, c.fill = negrita, fondo
        totales = {}
        for (k, mon), d in sorted(filas_dict.items()):
            prom = d["bruto"] / d["noches"] if d["noches"] else 0
            ws.append([k, mon, d["reservas"], d["noches"], round(d["bruto"], 2), round(d["comision"], 2), round(d["neto"], 2), round(prom, 2)])
            t = totales.setdefault(mon, {"reservas": 0, "noches": 0, "bruto": 0.0, "comision": 0.0, "neto": 0.0})
            for x in t:
                t[x] += d[x]
        for mon, t in sorted(totales.items()):
            prom = t["bruto"] / t["noches"] if t["noches"] else 0
            ws.append([f"TOTAL {anio}", mon, t["reservas"], t["noches"], round(t["bruto"], 2), round(t["comision"], 2), round(t["neto"], 2), round(prom, 2)])
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
        for col, ancho in zip("ABCDEFGH", (28, 9, 10, 9, 16, 14, 16, 18)):
            ws.column_dimensions[col].width = ancho
        for fila in ws.iter_rows(min_row=5):
            for c in fila[4:]:
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")

    hoja(wb.active, f"Resumen {anio} por mes", agrupar(lambda r: (r.get("fecha_inicio") or "")[:7]), "Mes")
    wb.active.title = "Por mes"
    hoja(wb.create_sheet("Por propiedad"), f"Resumen {anio} por propiedad",
         agrupar(lambda r: r.get("propiedad") or "Sin propiedad"), "Propiedad")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _escribir_codigo(z):
    """Copia el código del sistema (lista blanca de tipos de archivo, sin claves)."""
    n = 0
    for base, carpetas, archivos in os.walk(RAIZ):
        carpetas[:] = [c for c in carpetas if c not in CARPETAS_EXCLUIDAS]
        for nombre in archivos:
            ext = os.path.splitext(nombre)[1].lower()
            if nombre.endswith(".env") or nombre.upper().startswith("CLAVES") or nombre == ".env":
                continue
            if ext in EXTENSIONES_CODIGO or nombre in NOMBRES_CODIGO:
                ruta = os.path.join(base, nombre)
                z.write(ruta, "codigo/" + os.path.relpath(ruta, RAIZ).replace(os.sep, "/"))
                n += 1
    return n


def _leeme(anio, datos, n_codigo):
    r = datos["reservas"]
    return f"""ARCHIVO ANUAL {anio} — PalermitanHost
Generado el {hoy().strftime('%d/%m/%Y')}. Desarrollado por Zynnos.

GUARDÁ ESTE ARCHIVO EN UN LUGAR SEGURO (tu computadora y una copia en Drive o un pendrive).
Tiene datos personales de huéspedes: no lo compartas.

QUÉ HAY ADENTRO
1. Resumen_{anio}.xlsx
   Resumen del año por mes y por propiedad: reservas, noches, ingresos y precio promedio.
   Separado por moneda (pesos y dólares). Las reservas canceladas no se cuentan.

2. datos_{anio}/  (se abren con Excel)
   reservas.csv ............ {len(r)} reservas que empezaron en {anio}
   resenas.csv ............. {len(datos['resenas'])} reseñas
   consultas_huespedes.csv . {len(datos['consultas_huespedes'])} consultas de huéspedes
   informes.csv ............ {len(datos['informes'])} informes generados con IA
   propiedades.csv ......... tus propiedades y sus datos

3. base_completa/
   TODAS las tablas del sistema (todos los años) en CSV. Con esto se puede reconstruir el sistema.

4. codigo/  ({n_codigo} archivos)
   El programa completo. NO tiene claves: las claves (IA, base de datos, búsquedas) están
   guardadas solo en Render, en "Environment".

LOS DATOS SIGUEN EN EL SISTEMA
Descargar este archivo no borra nada. El sistema necesita los años anteriores para comparar
temporadas, calcular el precio sugerido y armar los informes.

SI ALGÚN DÍA EL SISTEMA SE ROMPE O SE PIERDEN DATOS
- Primero avisá a Zynnos (Pedro Santos).
- El código también está en tu repositorio de GitHub "palermitanhost".
- Copias de la base: GitHub → tu repositorio → Actions (semanales, 90 días)
  y Releases ("Copia anual"), cifradas con tu frase CLAVE_RESPALDO.
- Pasos técnicos para un programador:
  a) Crear una base PostgreSQL y un servicio web (Render) con el código de "codigo/".
     Instalar: pip install -r requirements.txt · Iniciar: gunicorn app:app
  b) Cargar las variables de entorno (ver codigo/.env.example).
  c) Al arrancar, el sistema crea las tablas solo.
  d) Restaurar datos: desde la copia cifrada (pg_restore, ver codigo/.github/workflows/respaldo.yml)
     o importando los CSV de base_completa/ con \\copy de PostgreSQL, tabla por tabla
     (primero propiedades, después el resto).
"""


def generar_zip(anio):
    datos = _datos_del_anio(anio)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"Resumen_{anio}.xlsx", _resumen_excel(anio, datos["reservas"]))
        for nombre, filas in datos.items():
            z.writestr(f"datos_{anio}/{nombre}.csv", _csv(filas))
        escribir_base_completa(z, "base_completa/")
        n = _escribir_codigo(z)
        z.writestr("LEEME.txt", _leeme(anio, datos, n))
    buf.seek(0)
    return buf
