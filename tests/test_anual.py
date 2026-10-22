# ### CAMBIO AQUÍ ### PRUEBA AUTOMÁTICA — ¡BORRA LA BASE QUE USA!
# Solo corre con PRUEBAS_OK=1 y una base local (PRUEBAS_DATABASE_URL en localhost). Nunca la de Render o Neon.
import os as _os, sys as _sys
from urllib.parse import urlsplit as _us
_U = _os.environ.get("PRUEBAS_DATABASE_URL", "")
if not _U or (_us(_U).hostname or "") not in ("", "localhost", "127.0.0.1") or _os.environ.get("PRUEBAS_OK") != "1":
    _sys.exit("Estas pruebas BORRAN la base: solo corren con PRUEBAS_OK=1 y PRUEBAS_DATABASE_URL apuntando a una base local.")
import os, sys, io, zipfile, warnings
from datetime import date
warnings.filterwarnings("ignore"); sys.path.insert(0, sys.argv[1])
os.environ.update(DATABASE_URL=_U, GROQ_API_KEY="x",
                  SECRET_KEY="clave-de-prueba-larga-123", ADMIN_USER="adriano", ADMIN_PASS="prueba-segura-123", TAREAS_TOKEN="z"*20)
import psycopg2
c0 = psycopg2.connect(os.environ["DATABASE_URL"]); c0.autocommit = True
c0.cursor().execute("DROP SCHEMA public CASCADE; CREATE SCHEMA public;"); c0.close()
import app as A, database as db, archivo_anual as aa
from openpyxl import load_workbook
fallas = 0
def chk(n, c, x=""):
    global fallas
    if not c: fallas += 1
    print(("OK " if c else "XX ") + n, "" if c else x)
# archivos trampa que NUNCA deben entrar al zip
R = sys.argv[1]
# ### CAMBIO AQUÍ ### nunca se pisa ni se borra un .env real: solo se crean (y después se borran) los que no existían
SECRETO1, SECRETO2 = "gsk_" + "secreto", "ADMIN_PASS=" + "secreto"
CREADOS = []
for _n, _t in (("CLAVES_TRAMPA.env", "GROQ_API_KEY=" + SECRETO1 + "\n"), (".env", SECRETO2 + "\n")):
    if not os.path.exists(os.path.join(R, _n)):
        open(os.path.join(R, _n), "w").write(_t); CREADOS.append(_n)
try:
    p1 = db.execute("INSERT INTO propiedades (nombre,barrio,moneda) VALUES ('Depto Soho','Palermo','USD') RETURNING id")
    p2 = db.execute("INSERT INTO propiedades (nombre,barrio,moneda) VALUES ('Studio VC','Villa Crespo','ARS') RETURNING id")
    filas = [(p1, "A1", "2025-01-10", 3, 300, 30, 270, "completada", "USD"), (p1, "A2", "2025-01-20", 2, 200, 20, 180, "completada", "USD"),
             (p1, "A3", "2025-02-05", 4, 400, 40, 360, "cancelada", "USD"), (p2, "A4", "2025-03-01", 2, 200000, 6000, 194000, "completada", "ARS"),
             (p1, "B1", "2026-05-01", 2, 260, 26, 234, "completada", "USD")]
    for pid, cod, f, n, b, c, ne, est, mon in filas:
        db.execute("""INSERT INTO reservas (propiedad_id,codigo_confirmacion,huesped_nombre,fecha_inicio,noches,ingreso_bruto,
                      comision_airbnb,ingreso_neto,estado,moneda) VALUES (%s,%s,'Lucía Pérez',%s,%s,%s,%s,%s,%s,%s)""", (pid, cod, f, n, b, c, ne, est, mon))
    db.execute("INSERT INTO resenas (propiedad_id,huesped,puntaje,texto,fecha) VALUES (%s,'X',5,'Genial','2025-01-15')", (p1,))
    db.execute("INSERT INTO resenas (propiedad_id,huesped,puntaje,texto,fecha) VALUES (%s,'Y',4,'Bien','2026-05-05')", (p1,))
    c = A.app.test_client()
    chk("pide login", c.get("/archivo-anual/2025").status_code == 302)
    c.post("/login", data={"usuario": "adriano", "password": "prueba-segura-123"})
    chk("años con datos", aa.anios_con_datos() == [2026, 2025], aa.anios_con_datos())
    chk("pendiente: el año terminado (2025)", aa.anio_pendiente() == 2025)
    h = c.get("/").get_data(as_text=True)
    chk("aviso fijo en el Dashboard", "Archivo anual 2025 sin descargar" in h)
    h = c.get("/informes").get_data(as_text=True)
    chk("Informes: botones por año", "2025 · pendiente" in h and "2026 (en curso)" in h, "")
    total_antes = db.query_one("SELECT COUNT(*) n FROM reservas")["n"]
    r = c.get("/archivo-anual/2025")
    z = zipfile.ZipFile(io.BytesIO(r.data)); nombres = z.namelist()
    chk("descarga zip", r.status_code == 200 and "Archivo-anual-2025" in r.headers["Content-Disposition"])
    chk("estructura del zip", all(x in nombres for x in ["Resumen_2025.xlsx", "datos_2025/reservas.csv", "datos_2025/resenas.csv",
        "datos_2025/consultas_huespedes.csv", "datos_2025/informes.csv", "datos_2025/propiedades.csv", "base_completa/reservas.csv",
        "base_completa/propiedades.csv", "codigo/app.py", "codigo/archivo_anual.py", "codigo/templates/dashboard.html",
        "codigo/requirements.txt", "codigo/.python-version", "codigo/.env.example", "codigo/.github/workflows/respaldo.yml", "LEEME.txt"]),
        [n for n in nombres if not n.startswith(("codigo/templates", "codigo/static"))][:40])
    chk("SIN claves: ni .env ni CLAVES_*.env", not any(n.endswith("/.env") or "CLAVES" in n for n in nombres)
        and all(SECRETO1.encode() not in z.read(n) and SECRETO2.encode() not in z.read(n) for n in nombres))
    res = z.read("datos_2025/reservas.csv").decode("utf-8-sig")
    import csv as _csv
    codigos = sorted(r["codigo_confirmacion"] for r in _csv.DictReader(io.StringIO(res)))
    chk("datos_2025 solo trae 2025 (4 reservas, con la cancelada)", codigos == ["A1", "A2", "A3", "A4"], codigos)
    chk("reseñas solo de 2025", "Genial" in z.read("datos_2025/resenas.csv").decode("utf-8-sig") and "Bien" not in z.read("datos_2025/resenas.csv").decode("utf-8-sig"))
    chk("base completa trae todos los años", "B1" in z.read("base_completa/reservas.csv").decode("utf-8-sig"))
    wb = load_workbook(io.BytesIO(z.read("Resumen_2025.xlsx")))
    filas_mes = [[c.value for c in r] for r in wb["Por mes"].iter_rows(min_row=5)]
    tot = {f[1]: f for f in filas_mes if f[0] == "TOTAL 2025"}
    chk("resumen: cancelada no suma reservas ni noches; lo cobrado sí (2, 5, 900)", tot["USD"][2:5] == [2, 5, 900], tot.get("USD"))
    chk("resumen: total ARS separado", tot["ARS"][2:5] == [1, 2, 200000], tot.get("ARS"))
    chk("resumen: bruto por noche USD = 180", tot["USD"][7] == 180)
    chk("resumen: hoja por propiedad", "Depto Soho" in [r[0].value for r in wb["Por propiedad"].iter_rows(min_row=5)])
    leeme = z.read("LEEME.txt").decode()
    chk("LEEME explica contenido y restauración", "4 reservas" in leeme and "pg_restore" in leeme and "no borra nada" in leeme.lower().replace("descargar este archivo no borra nada", "no borra nada"))
    chk("se marca descargado y el aviso se va", aa.descargado(2025) and aa.anio_pendiente() is None and "archivo anual 2025" not in c.get("/").get_data(as_text=True))
    chk("NO se borró ningún dato", db.query_one("SELECT COUNT(*) n FROM reservas")["n"] == total_antes)
    c.get("/archivo-anual/2026")
    chk("año en curso: se descarga pero no queda marcado", not aa.descargado(2026))
    chk("año sin datos → vuelve a Informes", c.get("/archivo-anual/2019").status_code == 302)
    aa.hoy = lambda: date(2027, 1, 1)
    chk("simulando 1/1/2027: pendiente 2026", aa.anio_pendiente() == 2026 and "Archivo anual 2026" in c.get("/").get_data(as_text=True))
    chk("Descargar copia (respaldo) sigue andando", "reservas.csv" in zipfile.ZipFile(io.BytesIO(c.get("/respaldo/descargar").data)).namelist())
finally:
    for _n in CREADOS: os.remove(os.path.join(R, _n))
print("\nFALLAS:", fallas)
